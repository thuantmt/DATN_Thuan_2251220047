"""Build the Week 3 report from saved verification evidence, without secrets."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from build_phase1_report import configure_styles, add_table

ROOT = Path(__file__).resolve().parents[2]
evidence = json.loads((ROOT / 'results/week3-verification.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
configure_styles(doc)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.add_paragraph('Nền tảng xác thực IELTS Writing', style='Title')
doc.add_paragraph('Báo cáo quy trình W03 01    Cập nhật thực tế ngày 14 tháng 10 năm 2026')
doc.add_paragraph(
    'Đã triển khai backend xác thực, phân quyền ba vai trò và frontend tài khoản. '
    'Các kiểm tra tự động đạt trên MongoDB cô lập. Tuần 3 chưa nghiệm thu đầy đủ vì '
    'kết nối MongoDB Atlas thật thất bại, chưa kiểm tra giao diện trên trình duyệt và chưa triển khai trực tuyến.'
)
doc.add_heading('1 Mục tiêu và đầu vào', level=1)
doc.add_paragraph(
    'Đồ án DATN_Thuan_2251220047 thực hiện từ 14/09/2026 đến 06/12/2026. '
    'Tuần 3 là nhãn sprint; ngày trên báo cáo là ngày thực hiện thực tế. '
    'Phạm vi được mở rộng từ frontend sang Express TypeScript, MongoDB, đăng ký, đăng nhập, phiên và RBAC.'
)
doc.add_paragraph(
    'Đã đọc AGENTS.md, README, requirements, decisions và reporting convention; kiểm tra Git và cấu hình. '
    'Ban đầu Git không có thay đổi chưa commit. Backend đã có nền hệ thống, React đã có các màn hình demo; '
    'chưa có xác thực thật. Giữ nguyên bộ đề, ảnh, prompt, schema và demo HTML.'
)
doc.add_heading('2 Công việc và quyết định', level=1)
add_table(doc, ['Hạng mục', 'Kết quả thực hiện'], [
    ('MongoDB', 'User và Session; email unique; session hash và TTL; seed ba vai trò chạy lại an toàn.'),
    ('Xác thực', 'Opaque session và cookie HttpOnly; mật khẩu scrypt; khôi phục phiên qua auth me; logout thu hồi server.'),
    ('Phân quyền', 'Middleware đọc role và status hiện tại. Dashboard chỉ dành đúng vai trò; tài khoản blocked bị từ chối.'),
    ('Frontend', 'Login, Register, API client, trạng thái phiên, dashboard và 403. Demo Writing riêng tại /demo.'),
    ('Vận hành', 'CORS, kiểm tra Origin và custom header chống CSRF, rate limit, cấu hình HTTPS, Vercel và Render.'),
], [1.3, 5.3])
doc.add_paragraph(
    'Chọn phiên lưu MongoDB để duy trì đăng nhập sau refresh và thu hồi ngay khi logout, không cần JWT secret. '
    'Không triển khai bài viết trên server, gọi AI, band score hoặc thống kê học tập thật. '
    'Dashboard chỉ hiển thị thông tin tài khoản được xác minh từ backend.'
)
doc.add_page_break()
doc.add_heading('3 Kiểm tra và bằng chứng', level=1)
add_table(doc, ['Kiểm tra', 'Kết quả', 'Bằng chứng trong results'], [
    ('Backend lint và type check', 'PASS', 'week3-backend-lint.txt và week3-backend-typecheck.txt'),
    ('Backend test', f"PASS {evidence['backend']['tests']} test", 'week3-backend-tests.txt'),
    ('Backend build và audit', 'PASS 0 lỗ hổng', 'week3-backend-build.txt và week3-backend-audit.json'),
    ('React test và type check', f"PASS {evidence['frontend']['tests']} test", 'week3-frontend-tests.txt và week3-frontend-typecheck.txt'),
    ('React production build', 'PASS', 'week3-frontend-build.txt'),
    ('Demo cũ và hợp đồng Python', 'PASS 7 và 26 test', 'week3-demo-tests.txt và week3-python-tests.txt'),
    ('Atlas và trình duyệt thật', 'BLOCKED', 'week3-verification.json'),
], [1.8, 1.25, 3.55])
doc.add_paragraph(
    'Test backend gọi Express qua HTTP và dùng MongoMemoryServer cô lập, không dùng Atlas hoặc database production. '
    'Đã kiểm tra đăng ký trùng đồng thời, ba vai trò, phiên sai và hết hạn, blocked, đổi quyền, CSRF, rate limit, '
    'cookie production, logout và seed lặp. Test React giả lập fetch; chưa phải kiểm thử trình duyệt với API thật.'
)
doc.add_heading('4 Hạn chế và phụ thuộc', level=1)
doc.add_paragraph(
    'DNS hệ thống trả ECONNREFUSED cho SRV. Resolver chỉ trong process với 1.1.1.1 và 8.8.8.8 phân giải được '
    'ba bản ghi; thử kết nối vẫn gặp MongoNetworkError và ERR_SSL_TLSV1_ALERT_INTERNAL_ERROR. '
    'Chưa xác định nguyên nhân cuối cùng. Cần kiểm tra Atlas Network Access, IP allowlist, trạng thái cluster và mạng TLS. '
    'Không đổi DNS hệ điều hành hoặc tắt xác minh TLS.'
)
doc.add_paragraph(
    'Không có trình duyệt kết nối nên chưa nghiệm thu trực quan desktop và mobile. '
    'Chưa deploy, chưa kiểm tra persistence qua restart production. Rate limit hiện ở bộ nhớ process; '
    'nhiều instance cần shared store. Cookie cross-site có thể bị trình duyệt chặn; ưu tiên domain app và api cùng site.'
)
doc.add_heading('5 Vận hành và điều kiện nghiệm thu', level=1)
doc.add_paragraph(
    'Hướng dẫn VS Code, biến môi trường, seed và triển khai nằm trong docs/week3-deployment.md; '
    'thiết kế API nằm trong backend/docs/authentication.md; danh sách tệp nằm trong results/week3-files.json. '
    'Chạy backend npm run dev và frontend/react npm run dev sau khi cấu hình MongoDB. '
    'Đăng ký công khai chỉ tạo learner; teacher và admin dùng seed trong môi trường development với mật khẩu từ biến môi trường.'
)
doc.add_paragraph(
    'Nghiệm thu còn cần ba tài khoản đăng nhập trên Atlas qua frontend thật, refresh giữ phiên, thử sai quyền bị từ chối, '
    'logout vô hiệu phiên cũ và dữ liệu còn sau restart. Mốc triển khai trực tuyến chỉ hoàn thành sau smoke test HTTPS. '
    'Chưa có AI thật hoặc bằng chứng độ chính xác IELTS.'
)
output = ROOT / 'docs/W03_01_Nen_tang_xac_thuc.docx'
doc.save(output)
print(output)
