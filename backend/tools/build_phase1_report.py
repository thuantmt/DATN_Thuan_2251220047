from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs" / "W02_01_Khoi_tao_backend.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def add_table(doc, headers, rows, widths) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, (header, width) in enumerate(zip(headers, widths)):
        cell = header_row.cells[index]
        cell.width = Inches(width)
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[index]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F6FB")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def configure_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        if style_name == "Title" and style._element.pPr is not None:
            border = style._element.pPr.find(qn("w:pBdr"))
            if border is not None:
                style._element.pPr.remove(border)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Khởi tạo nền tảng backend IELTS Writing")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Báo cáo quy trình W02 01   Ngày cập nhật 23 tháng 09 năm 2026")
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(75, 85, 99)

    doc.add_heading("Kết luận thực hiện", level=1)
    doc.add_paragraph(
        "Giai đoạn 1 đã tạo được nền Express và TypeScript, kiểm tra cấu hình bằng Zod, lớp kết nối "
        "Mongoose, health và readiness, request ID, logging an toàn, error envelope và nền validation. "
        "Lint, type check, test tự động và build đều đạt. Kết nối Atlas thật chưa được nghiệm thu vì lần "
        "chạy ngày 23 tháng 09 năm 2026 trả ECONNREFUSED trước khi server mở cổng; do đó báo cáo không "
        "xác nhận database, JWT hoặc Gemini đã hoạt động."
    )

    doc.add_heading("1 Mục tiêu và phạm vi", level=1)
    doc.add_paragraph(
        "Quy trình này thiết lập nền kỹ thuật để các giai đoạn sau bổ sung xác thực JWT, đề và nháp, "
        "phiên bản bài viết, Gemini, giáo viên và quản trị. Phạm vi chỉ gồm Giai đoạn 1 theo prompt đã "
        "được phê duyệt; frontend, dữ liệu mẫu, prompt và JSON Schema không bị thay đổi."
    )
    add_bullet(doc, "Backend dùng Node.js 24, TypeScript, Express, Mongoose, Zod, Helmet, CORS và Pino.")
    add_bullet(doc, "Server chỉ lắng nghe trên 127.0.0.1 sau khi MongoDB kết nối thành công.")
    add_bullet(doc, "Test HTTP dùng database probe giả lập và không truy cập Atlas production.")

    doc.add_heading("2 Công việc đã thực hiện", level=1)
    add_table(
        doc,
        ["Hạng mục", "Kết quả", "Bằng chứng chính"],
        [
            ("Nền dự án", "Đã tạo package, TypeScript, ESLint, Vitest và build scripts", "backend/package.json"),
            ("Cấu hình", "Zod kiểm tra Atlas URI, port, origin và cấu hình tương lai", "src/config/environment.ts"),
            ("Database", "Kết nối Mongoose có timeout, tắt buffer và readiness probe", "src/database/mongoose.ts"),
            ("HTTP", "Request ID, logging tối thiểu, Helmet, CORS và JSON limit", "src/app.ts và src/middleware"),
            ("Endpoint", "Đã có health và ready dưới base path api v1", "src/routes/system.ts"),
            ("Hợp đồng", "Đã mô tả 10 collection, API JWT và EvaluationProvider", "backend/docs và src/contracts"),
        ],
        [1.35, 3.15, 1.95],
    )

    doc.add_heading("3 Kiến trúc và endpoint", level=1)
    doc.add_paragraph(
        "Luồng khởi động đọc và kiểm tra biến môi trường, kết nối Atlas, sau đó mới tạo HTTP listener. "
        "Luồng request tạo ID mới tại server, ghi log không có body hoặc secret, áp dụng middleware bảo "
        "mật và trả lỗi thống nhất. Những endpoint chưa triển khai trả 404 thay vì dữ liệu mô phỏng."
    )
    add_table(
        doc,
        ["Phương thức", "Đường dẫn", "Kết quả"],
        [
            ("GET", "/api/v1/health", "200 khi tiến trình Node.js hoạt động"),
            ("GET", "/api/v1/ready", "200 khi MongoDB connected; 503 nếu chưa sẵn sàng"),
        ],
        [1.15, 2.35, 3.0],
    )

    doc.add_heading("4 Kết quả kiểm tra", level=1)
    add_table(
        doc,
        ["Kiểm tra", "Kết quả", "Ghi chú"],
        [
            ("npm run lint", "Đạt", "Không còn lỗi ESLint"),
            ("npm run type-check", "Đạt", "TypeScript strict không báo lỗi"),
            ("npm test", "Đạt", "3 tệp test và 7 test backend"),
            ("npm run build", "Đạt", "Sinh mã JavaScript trong dist"),
            ("npm audit", "Đạt tại lúc cài", "npm báo 0 vulnerability"),
            ("Atlas thật", "Chưa đạt", "Kết nối trả ECONNREFUSED; server không mở cổng"),
        ],
        [1.65, 1.15, 3.7],
    )

    doc.add_heading("5 Hạn chế và phụ thuộc", level=1)
    add_bullet(doc, "Cần kiểm tra Atlas Network Access, IP allowlist, firewall và khả năng mở outbound MongoDB từ máy chạy.")
    add_bullet(doc, "Chưa có collection hoặc seed; tài liệu collection mới là thiết kế cho giai đoạn sau.")
    add_bullet(doc, "Chưa triển khai JWT, refresh token, phân quyền và ownership.")
    add_bullet(doc, "Chưa gọi Gemini, chưa có đánh giá hoặc band score thật.")
    add_bullet(doc, "Frontend vẫn sử dụng localStorage và phản hồi có nhãn mô phỏng.")

    doc.add_heading("6 Việc tiếp theo và điều kiện nghiệm thu", level=1)
    doc.add_paragraph(
        "Trước khi mở Giai đoạn 2, cần cho phép IP hiện tại trong Atlas, xác nhận database user có quyền "
        "readWrite trên database ielts_writing và kiểm tra mạng không chặn kết nối MongoDB. Chạy lại "
        "backend và yêu cầu cả health lẫn ready trả 200. Sau khi điều kiện này đạt, Giai đoạn 2 mới triển "
        "khai user, refresh session, access token, refresh rotation, logout và kiểm tra 401, 403."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
